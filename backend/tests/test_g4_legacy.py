"""Chehreh / Tizpardaz file imports: big files are refused (never cut), odd Excel sizes and UTF-16 text are read,
undo keeps deposits that were used, a re-import never brings a used one back, automatic codes never catch another person."""
import inspect
import io
import time
import zipfile

import pytest

from app.services import legacy_import as li


def _xlsx(rows):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _post(client, name, data, kind, unit="rial"):
    return client.post("/api/import/legacy/preview", files={"file": (name, data, "application/octet-stream")},
                       data={"kind": kind, "unit": unit})


def _customer(client, mobile):
    items = client.get(f"/api/customers?q={mobile}").json()["items"]
    assert len(items) == 1, items
    return items[0]


def _balance(client, cid):
    return client.get(f"/api/customers/{cid}").json()["balance"]


def _tb(client):
    t = client.get("/api/ledger/trial-balance").json()
    assert t["total_debit"] == t["total_credit"]
    return {r["code"]: r["balance"] for r in t["rows"]}


# ------------------------------------------------------------------ too many rows: a clear message, nothing cut
def test_more_rows_than_the_limit_is_refused_not_cut(client, monkeypatch):
    head = "نام مشتری,موبایل\n"
    rows = "".join(f"مشتری حد {i},0915104{i:04d}\n" for i in range(1, 50_001))  # header + 50,000 = 50,001 filled rows
    r = _post(client, "big.csv", (head + rows).encode(), "customers")
    assert r.status_code == 400 and "۵۰٬۰۰۰ ردیف" in r.json()["detail"] and "جدا" in r.json()["detail"]

    monkeypatch.setattr(li, "MAX_ROWS", 4)
    ok = [["نام مشتری", "موبایل"], *([f"حد کم {i}", f"0915104900{i}"] for i in range(1, 4))]  # 4 rows: all read
    assert len(li.read_table("a.xlsx", _xlsx(ok))) == 4
    assert len(li.read_table("a.csv", "\n".join(",".join(r) for r in ok).encode())) == 4
    too_many = [*ok, ["حد کم 4", "09151049004"]]
    for name, data in (("b.xlsx", _xlsx(too_many)), ("b.csv", "\n".join(",".join(r) for r in too_many).encode()),
                       ("b.xls", ("<table>" + "".join(f"<tr><td>{a}</td><td>{b}</td></tr>" for a, b in too_many)
                                  + "</table>").encode())):
        with pytest.raises(li.ImportProblem, match="ردیف"):
            li.read_table(name, data)
    # the same for the Tizpardaz import (it reads files the same way)
    r = client.post("/api/import/tizpardaz/preview", files={"file": ("tp.xlsx", _xlsx(too_many), "application/octet-stream")},
                    data={"kind": "customers"})
    assert r.status_code == 400 and "ردیف" in r.json()["detail"]
    # rows handed straight to parse_rows are checked too
    from app.core.db import SessionLocal
    with SessionLocal() as db, pytest.raises(li.ImportProblem, match="ردیف"):
        li.parse_rows(db, too_many + [["حد کم 5", "09151049005"]], "customers", "rial")


# ------------------------------------------------------------------ a sheet that declares a huge size
def test_xlsx_with_a_bogus_huge_size_is_read_quickly(client):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["نام مشتری", "موبایل"])
    ws.append(["ابعاد بزرگ", "09151040301"])
    ws.cell(row=2, column=16384, value="دور")  # a stray cell in the last column (XFD)
    ws.cell(row=1_048_576, column=1, value="آخر")  # and one in the very last row
    buf = io.BytesIO()
    wb.save(buf)
    # make the sheet declare the largest possible size, as a broken / hostile file can
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                start = data.index(b'<dimension ref="')
                end = data.index(b'"', start + 16)
                data = data[:start] + b'<dimension ref="A1:XFD1048576' + data[end:]
            z.writestr(item, data)
    t0 = time.monotonic()
    rows = li.read_table("big.xlsx", out.getvalue())
    assert time.monotonic() - t0 < 15  # it used to take minutes (and froze the program)
    assert [r[:2] for r in rows] == [["نام مشتری", "موبایل"], ["ابعاد بزرگ", "09151040301"], ["آخر", None]]
    assert {len(r) for r in rows} == {2}  # the stray far cell is ignored, every row has the table's width
    r = _post(client, "big.xlsx", out.getvalue(), "customers")
    assert r.status_code == 200, r.text
    assert r.json()["summary"]["ok"] == 2


def test_xlsx_rows_with_a_far_right_cell_are_never_built_full_width(monkeypatch):
    """(review) Thousands of rows that each hold one (styled, empty) cell in column XFD: every row used to be built
    16,384 cells wide (minutes for a small hostile file). Now no row is read wider than MAX_COLS."""
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet

    head = ('<row r="1"><c r="A1" t="inlineStr"><is><t>نام مشتری</t></is></c><c r="B1" t="inlineStr"><is><t>موبایل</t></is></c></row>'
            '<row r="2"><c r="A2" t="inlineStr"><is><t>ستون دور</t></is></c><c r="B2" t="inlineStr"><is><t>09151040501</t></is></c>'
            '<c r="XFD2" s="1"/></row>')
    far = "".join(f'<row r="{i}"><c r="XFD{i}" s="1"/></row>' for i in range(3, 3003))
    sheet = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
             f'<dimension ref="A1:XFD3002"/><sheetData>{head}{far}</sheetData></worksheet>')
    buf = io.BytesIO(_xlsx([["x"]]))
    src = zipfile.ZipFile(buf)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for item in src.infolist():
            z.writestr(item, sheet.encode() if item.filename == "xl/worksheets/sheet1.xml" else src.read(item.filename))
    widest = [0]
    real = ReadOnlyWorksheet._get_row

    def spy(self, *a, **kw):  # noqa: ANN001, ANN002, ANN003
        r = real(self, *a, **kw)
        widest[0] = max(widest[0], len(r))
        return r
    monkeypatch.setattr(ReadOnlyWorksheet, "_get_row", spy)
    rows = li.read_table("far.xlsx", out.getvalue())
    assert rows == [["نام مشتری", "موبایل"], ["ستون دور", "09151040501"]]
    assert 0 < widest[0] <= li.MAX_COLS


def test_upload_handlers_do_not_block_the_server():
    from app.api import legacy, tizpardaz
    assert not inspect.iscoroutinefunction(legacy.preview)
    assert not inspect.iscoroutinefunction(tizpardaz.preview)


# ------------------------------------------------------------------ text encodings
def test_utf16_text_exports_are_read(client):
    text = "نام مشتری\tموبایل\nسارا یونیکد\t09151040401\n"
    want = [["نام مشتری", "موبایل"], ["سارا یونیکد", "09151040401"]]
    for data in (text.encode("utf-16"),  # Excel «Unicode Text»: UTF-16 with its mark
                 text.encode("utf-16-le"), text.encode("utf-16-be"),  # without the mark
                 text.encode("utf-8-sig"), text.encode("utf-8")):
        assert li.read_table("x.txt", data) == want, data[:10]
    arabic = text.replace("ی", "ي")  # Windows Arabic code page (old programs): still read as before
    assert li.read_table("x.csv", arabic.encode("cp1256")) == [[c.replace("ی", "ي") for c in r] for r in want]
    r = _post(client, "unicode.txt", text.encode("utf-16"), "customers")
    assert r.status_code == 200, r.text
    assert r.json()["summary"]["ok"] == 1 and r.json()["sample"][0]["mobile"] == "09151040401"
    # an HTML table saved as UTF-16 with an .xls name
    html = "<html><table><tr><td>نام مشتری</td><td>موبایل</td></tr><tr><td>سارا یونیکد</td><td>09151040401</td></tr></table></html>"
    assert li.read_table("r.xls", html.encode("utf-16")) == want


# ------------------------------------------------------------------ undo keeps a used deposit
def test_undo_keeps_a_deposit_used_on_a_voided_invoice(client, services):
    svc = services["کاشت ناخن"]
    before = _tb(client)
    dep_file = _xlsx([["نام مشتری", "موبایل", "مبلغ بیعانه"],
                      ["بیعانه برگشتی جی‌چهار", "09151040501", 2_000_000],
                      ["بیعانه نصفه جی‌چهار", "09151040502", 3_000_000]])
    pv = _post(client, "dep.xlsx", dep_file, "deposits").json()
    assert client.post(f"/api/import/legacy/{pv['id']}/commit", json={}).json()["result"]["deposits"] == 2
    a, b = _customer(client, "09151040501"), _customer(client, "09151040502")
    da = client.get(f"/api/deposits?customer_id={a['id']}").json()[0]
    db_ = client.get(f"/api/deposits?customer_id={b['id']}").json()[0]
    # A: the whole deposit pays a 2,000,000 invoice; B: 2,000,000 of the 3,000,000 deposit (it is split)
    for cust, dep in ((a, da), (b, db_)):
        inv = client.post("/api/invoices", json={"customer_id": cust["id"], "deposit_ids": [dep["id"]],
                                                 "items": [{"service_id": svc["id"], "unit_price": 2_000_000}]})
        assert inv.status_code == 200, inv.text
        assert inv.json()["status"] == "paid"
        assert client.post(f"/api/invoices/{inv.json()['id']}/void", params={"payments": "cancel"}).status_code == 200
    assert _balance(client, a["id"]) == {"receivable": 0, "deposits_held": 2_000_000, "net": 2_000_000}
    assert _balance(client, b["id"]) == {"receivable": 0, "deposits_held": 3_000_000, "net": 3_000_000}

    u = client.post(f"/api/import/legacy/{pv['id']}/undo").json()
    assert u["kept"]["deposits"] == 2 and u["removed"]["deposits"] == 0
    # nothing left half-deleted: the books still match the deposits, to the Rial
    assert _balance(client, a["id"]) == {"receivable": 0, "deposits_held": 2_000_000, "net": 2_000_000}
    assert _balance(client, b["id"]) == {"receivable": 0, "deposits_held": 3_000_000, "net": 3_000_000}
    held_a = client.get(f"/api/deposits?status=held&customer_id={a['id']}").json()
    held_b = client.get(f"/api/deposits?status=held&customer_id={b['id']}").json()
    assert [d["amount"] for d in held_a] == [2_000_000]
    assert sorted(d["amount"] for d in held_b) == [1_000_000, 2_000_000]
    after = _tb(client)
    assert after.get("2100", 0) - before.get("2100", 0) == 5_000_000  # still owed to the two customers
    assert after.get("3200", 0) - before.get("3200", 0) == -5_000_000  # their opening balances stay
    assert after.get("1200", 0) - before.get("1200", 0) == 0


def test_undo_still_removes_an_unused_deposit_completely(client):
    before = _tb(client)
    pv = _post(client, "dep2.xlsx", _xlsx([["نام مشتری", "موبایل", "مبلغ بیعانه"], ["بیعانه دست‌نخورده جی‌چهار", "09151040511", 700_000]]),
               "deposits").json()
    client.post(f"/api/import/legacy/{pv['id']}/commit", json={})
    u = client.post(f"/api/import/legacy/{pv['id']}/undo").json()
    assert u["removed"]["deposits"] == 1 and u["kept"]["deposits"] == 0 and u["removed"]["customers"] == 1
    after = _tb(client)
    assert after.get("2100", 0) == before.get("2100", 0) and after.get("3200", 0) == before.get("3200", 0)


# ------------------------------------------------------------------ re-import after undo
def test_reimport_after_undo_does_not_bring_back_a_used_deposit(client, services):
    svc = services["کاشت ناخن"]
    dep_file = _xlsx([["نام مشتری", "موبایل", "مبلغ بیعانه"],
                      ["نگار سلامی جی‌چهار", "09151040601", 1_000_000],
                      ["شیما سلامی جی‌چهار", "09151040602", 2_000_000]])
    pv = _post(client, "dep3.xlsx", dep_file, "deposits").json()
    client.post(f"/api/import/legacy/{pv['id']}/commit", json={})
    nagar = _customer(client, "09151040601")
    dep = client.get(f"/api/deposits?customer_id={nagar['id']}").json()[0]
    inv = client.post("/api/invoices", json={"customer_id": nagar["id"], "deposit_ids": [dep["id"]],
                                             "items": [{"service_id": svc["id"], "unit_price": 1_000_000}]})
    assert inv.status_code == 200, inv.text
    u = client.post(f"/api/import/legacy/{pv['id']}/undo").json()
    assert u["kept"]["deposits"] == 1 and u["removed"]["deposits"] == 1
    before = _tb(client)
    # the owner imports the same file again
    pv2 = _post(client, "dep3.xlsx", dep_file, "deposits").json()
    r = client.post(f"/api/import/legacy/{pv2['id']}/commit", json={}).json()["result"]
    assert r["deposits"] == 1 and r["skipped_duplicates"] == 1
    assert [(d["amount"], d["status"]) for d in client.get(f"/api/deposits?customer_id={nagar['id']}").json()] == [(1_000_000, "applied")]
    shima = _customer(client, "09151040602")
    assert [(d["amount"], d["status"]) for d in client.get(f"/api/deposits?customer_id={shima['id']}").json()] == [(2_000_000, "held")]
    assert _balance(client, nagar["id"])["deposits_held"] == 0
    assert _balance(client, shima["id"])["deposits_held"] == 2_000_000
    after = _tb(client)
    assert after.get("2100", 0) - before.get("2100", 0) == 2_000_000  # only Shima's deposit came back
    assert after.get("3200", 0) - before.get("3200", 0) == -2_000_000
    # the file is still remembered after a second undo / import round
    client.post(f"/api/import/legacy/{pv2['id']}/undo")
    pv3 = _post(client, "dep3.xlsx", dep_file, "deposits").json()
    r3 = client.post(f"/api/import/legacy/{pv3['id']}/commit", json={}).json()["result"]
    assert r3["deposits"] == 1 and r3["skipped_duplicates"] == 1


def test_reimport_after_undo_skips_receipts_that_were_kept(client, services):
    from datetime import datetime, timedelta

    from app.services.jalali import gregorian_to_jalali
    j = lambda d: "%04d/%02d/%02d" % gregorian_to_jalali(d.year, d.month, d.day)  # noqa: E731
    soon, past = datetime.now() + timedelta(days=5), datetime.now() - timedelta(days=30)
    hist = _xlsx([["نام مشتری", "موبایل", "تاریخ", "ساعت", "خدمت", "مبلغ"],
                  ["رسید آینده جی‌چهار", "09151040701", j(soon), "11:00", "کاشت ناخن", 900_000],
                  ["رسید آینده جی‌چهار", "09151040701", j(past), "12:00", "کاشت ناخن", 800_000]])
    pv = _post(client, "h.xlsx", hist, "history").json()
    assert client.post(f"/api/import/legacy/{pv['id']}/commit", json={}).json()["result"]["appointments"] == 2
    c = _customer(client, "09151040701")
    booked = client.get(f"/api/appointments?customer_id={c['id']}&status=booked").json()
    assert len(booked) == 1
    inv = client.post("/api/invoices", json={"customer_id": c["id"], "appointment_ids": [booked[0]["id"]],
                                             "items": [{"service_id": services["کاشت ناخن"]["id"], "unit_price": 900_000}]})
    assert inv.status_code == 200, inv.text
    u = client.post(f"/api/import/legacy/{pv['id']}/undo").json()
    assert u["kept"]["appointments"] == 1 and u["removed"]["appointments"] == 1
    pv2 = _post(client, "h.xlsx", hist, "history").json()
    r = client.post(f"/api/import/legacy/{pv2['id']}/commit", json={}).json()["result"]
    assert r["appointments"] == 1 and r["skipped_duplicates"] == 1


# ------------------------------------------------------------------ automatic customer codes
def test_an_automatic_code_is_not_matched_to_another_person(client):
    pv = _post(client, "c.xlsx", _xlsx([["نام مشتری", "موبایل"], ["الهام بی‌کد جی‌چهار", "09151040801"]]), "customers").json()
    client.post(f"/api/import/legacy/{pv['id']}/commit", json={})
    elham = _customer(client, "09151040801")
    auto = elham["code"]
    assert auto and auto.isdigit()
    # the old software later gives that same number to its next real customer
    dep = _xlsx([["کد مشتری", "نام مشتری", "موبایل", "مبلغ بیعانه"], [auto, "سمیرا جدید جی‌چهار", "09151040802", 5_000_000]])
    pv2 = _post(client, "d.xlsx", dep, "deposits").json()
    r = client.post(f"/api/import/legacy/{pv2['id']}/commit", json={}).json()["result"]
    assert r["customers_new"] == 1 and r["customers_updated"] == 0 and r.get("codes_moved") == 1
    samira = _customer(client, "09151040802")
    elham = _customer(client, "09151040801")
    assert samira["code"] == auto and elham["code"] != auto and elham["code"].isdigit()
    assert _balance(client, samira["id"])["deposits_held"] == 5_000_000
    assert _balance(client, elham["id"])["deposits_held"] == 0
    assert [d["amount"] for d in client.get(f"/api/deposits?customer_id={elham['id']}").json()] == []


def test_an_automatic_code_still_finds_the_same_person(client):
    pv = _post(client, "c2.xlsx", _xlsx([["نام مشتری", "موبایل"], ["نگار بی‌کد جی‌چهار", "09151040811"],
                                          ["پریسا بی‌کد جی‌چهار", "09151040812"]]), "customers").json()
    client.post(f"/api/import/legacy/{pv['id']}/commit", json={})
    nagar, parisa = _customer(client, "09151040811"), _customer(client, "09151040812")
    # Nagar: the old software has her under the very number given here -> same person, linked
    # Parisa: the old software has her under another (real) code -> she gets that code, no duplicate
    dep = _xlsx([["کد مشتری", "نام مشتری", "موبایل", "مبلغ بیعانه"],
                 [nagar["code"], "نگار بی‌کد جی‌چهار", "09151040811", 400_000],
                 ["G4P812", "پریسا بی‌کد جی‌چهار", "09151040812", 600_000]])
    pv2 = _post(client, "d2.xlsx", dep, "deposits").json()
    r = client.post(f"/api/import/legacy/{pv2['id']}/commit", json={}).json()["result"]
    assert r["customers_new"] == 0 and r["customers_updated"] == 2 and not r.get("codes_moved")
    assert _customer(client, "09151040811")["code"] == nagar["code"]
    assert _customer(client, "09151040812")["code"] == "G4P812"
    assert _balance(client, nagar["id"])["deposits_held"] == 400_000
    assert _balance(client, parisa["id"])["deposits_held"] == 600_000
    # Nagar's number is now confirmed by the old software: a receipt under it without a mobile still finds her
    from datetime import datetime, timedelta

    from app.services.jalali import gregorian_to_jalali
    d = datetime.now() - timedelta(days=3)
    hist = _xlsx([["کد مشتری", "نام مشتری", "تاریخ", "خدمت"],
                  [nagar["code"], "نگار ب. جی‌چهار", "%04d/%02d/%02d" % gregorian_to_jalali(d.year, d.month, d.day), "کاشت ناخن"]])
    pv3 = _post(client, "h3.xlsx", hist, "history").json()
    r3 = client.post(f"/api/import/legacy/{pv3['id']}/commit", json={}).json()["result"]
    assert r3["customers_new"] == 0 and not r3.get("codes_moved")
    assert len(client.get(f"/api/appointments?customer_id={nagar['id']}&status=done").json()) == 1


def test_undo_forgets_the_automatic_codes_of_the_customers_it_removed(client):
    """(review) A removed customer's record number can be given to the next customer (SQLite reuses it): its
    automatic code must not make that next person "automatic" - their real code would then be moved away."""
    from app.core.db import SessionLocal
    from app.models import ImportBatch

    pv = _post(client, "c3.xlsx", _xlsx([["نام مشتری", "موبایل"], ["ژیلا بی‌کد جی‌چهار", "09151040821"],
                                          ["مینو بی‌کد جی‌چهار", "09151040822"]]), "customers").json()
    client.post(f"/api/import/legacy/{pv['id']}/commit", json={})
    jila, minoo = _customer(client, "09151040821"), _customer(client, "09151040822")
    # Minoo pays a deposit here, so she stays when the file is taken back; Jila goes
    acc = client.get("/api/accounts").json()[0]["id"]
    r = client.post("/api/deposits", json={"customer_id": minoo["id"], "amount": 1_000_000, "payment_account_id": acc})
    assert r.status_code == 200, r.text
    assert client.post(f"/api/import/legacy/{pv['id']}/undo").status_code == 200
    assert client.get(f"/api/customers/{jila['id']}").status_code == 404
    with SessionLocal() as db:
        auto = db.get(ImportBatch, pv["id"]).summary["auto_codes"]
    assert auto == {str(minoo["id"]): minoo["code"]}  # Minoo's number is still one given here, Jila's is forgotten


def test_an_automatic_code_of_a_merged_duplicate_is_not_matched_to_another_person(client):
    pv = _post(client, "c5.xlsx", _xlsx([["نام مشتری", "موبایل"], ["مهسا ادغام جی‌چهار", "09151040901"],
                                          ["مهسا ادغام جی‌چهار", "09151040902"]]), "customers").json()
    client.post(f"/api/import/legacy/{pv['id']}/commit", json={})
    a, b = _customer(client, "09151040901"), _customer(client, "09151040902")
    assert client.post("/api/customers/merge", json={"keep_id": a["id"], "drop_ids": [b["id"]]}).status_code == 200
    kept = client.get(f"/api/customers/{a['id']}").json()
    assert b["code"] in (kept.get("other_codes") or [])
    # the old software later gives the dropped record's automatic number to someone else
    dep = _xlsx([["کد مشتری", "نام مشتری", "موبایل", "مبلغ بیعانه"], [b["code"], "رویا دیگری جی‌چهار", "09151040903", 700_000]])
    pv2 = _post(client, "d5.xlsx", dep, "deposits").json()
    r = client.post(f"/api/import/legacy/{pv2['id']}/commit", json={}).json()["result"]
    assert r["customers_new"] == 1
    roya = _customer(client, "09151040903")
    assert _balance(client, roya["id"])["deposits_held"] == 700_000
    assert _balance(client, a["id"])["deposits_held"] == 0
    kept = client.get(f"/api/customers/{a['id']}").json()
    assert kept["code"] == a["code"] and b["code"] not in (kept.get("other_codes") or [])
