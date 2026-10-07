"""Bring data over from the salon's previous software (e.g. «چهره») via its Excel / CSV exports.

Three kinds of file:
  * customers - name, mobile (+ birth date, notes)
  * history   - services already done: customer, date, service (+ staff, amount). Stored as completed
                appointments for the customer's history; NO money is posted (it was accounted in the old system).
  * deposits  - deposits held for future appointments: customer, amount, received date (+ service, appointment
                date). Posted as an opening balance (Dr opening equity / Cr customer deposits) - the cash itself
                is not counted again - and applied to the invoice like any other deposit later.

Each import is a batch: preview first (nothing written), then commit, and it can be undone completely.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime, time

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..models import (
    Appointment,
    Customer,
    Deposit,
    ImportBatch,
    Invoice,
    InvoiceItem,
    JournalEntry,
    Payment,
    PaymentAccount,
    Service,
    ServiceLine,
    Staff,
    WaitlistEntry,
    local_now,
)
from . import accounting, scheduling, service_catalog
from .jalali import jalali_to_gregorian
from .textutil import normalize_mobile, normalize_text, parse_amount, to_en_digits

KINDS = {"customers": "مشتریان", "history": "سوابق خدمات انجام‌شده", "deposits": "بیعانه‌های نوبت‌های آینده"}
IMPORT_NOTE = "انتقال از نرم‌افزار قبلی"
MAX_ROWS = 50_000

# column titles seen in Persian salon software exports (e.g. «چهره») -> field
SYNONYMS: dict[str, list[str]] = {
    "customer_code": ["کد مشتری", "شماره مشتری", "کد پرونده", "شماره پرونده", "کد"],
    "receipt_no": ["شماره پذیرش", "شماره فیش", "شماره فاکتور", "شماره رسید"],
    "name": ["نام مشتری", "مشتری", "نام و نام خانوادگی", "نام ونام خانوادگی", "نام کامل", "اسم", "مراجعه کننده", "نام مراجعه کننده",
             "خانم", "customer", "name"],
    "first_name": ["نام"],
    "last_name": ["نام خانوادگی", "فامیل", "فامیلی"],
    "mobile": ["موبایل", "تلفن همراه", "همراه", "شماره همراه", "شماره موبایل", "شماره تماس", "تلفن", "mobile", "phone"],
    "date": ["تاریخ", "تاریخ مراجعه", "تاریخ خدمت", "تاریخ انجام", "تاریخ دریافت", "تاریخ پرداخت", "تاریخ ثبت", "تاریخ فاکتور", "date"],
    "time": ["ساعت", "زمان", "ساعت مراجعه", "ساعت نوبت", "time"],
    "appt_date": ["تاریخ نوبت", "نوبت", "زمان نوبت", "تاریخ رزرو", "تاریخ مراجعه بعدی"],
    "service": ["خدمت", "خدمات", "نام خدمت", "عنوان خدمت", "شرح خدمت", "نوع خدمت", "شرح", "عنوان", "service"],
    "line": ["لاین", "بخش", "لاین خدمت", "دپارتمان", "واحد"],
    "staff": ["پرسنل", "آرایشگر", "کارشناس", "اپراتور", "نام پرسنل", "انجام دهنده", "متخصص", "کارمند", "staff"],
    "amount": ["قابل پرداخت", "مبلغ", "مبلغ بیعانه", "بیعانه", "پیش پرداخت", "مبلغ کل", "قیمت", "مبلغ پرداختی", "مبلغ دریافتی",
               "amount"],
    "status": ["وضعیت", "وضعیت بیعانه"],
    "settled_date": ["تاریخ تسویه"],
    "refund_date": ["تاریخ استرداد", "تاریخ برگشت"],
    "notes": ["توضیحات", "یادداشت", "ملاحظات", "توضیح", "notes"],
    "birth_date": ["تاریخ تولد", "تولد"],
}
# per file kind: deposits are paid on one date and the visit is another («تاریخ مراجعه» = appointment)
KIND_SYNONYMS: dict[str, dict[str, list[str]]] = {
    "deposits": {"date": ["تاریخ پرداخت", "تاریخ دریافت", "تاریخ واریز", "تاریخ", "تاریخ ثبت"],
                 "appt_date": ["تاریخ مراجعه", "تاریخ نوبت", "نوبت", "زمان نوبت", "تاریخ رزرو"],
                 "time": ["ساعت نوبت", "ساعت مراجعه"],
                 "amount": ["مبلغ بیعانه", "بیعانه", "پیش پرداخت", "مبلغ", "مبلغ دریافتی", "مبلغ پرداختی", "amount"]},
}
# report columns that are never imported (kept out so no other field grabs them by a partial match)
IGNORED = ["ردیف", "صندوقدار", "کد اشتراک", "کارتخوان", "دستیار", "درصد دستیار", "درصد", "سهم پرسنل", "تخفیف پرسنل", "تخفیف سالن",
           "غیرنقدی", "کدپیگیری", "کد پیگیری", "عروس", "پکیج", "مسترد کننده", "ساعت استرداد", "ویرایش کننده", "تاریخ ویرایش",
           "ساعت ویرایش", "ساعت ثبت", "شماره قرارگاه"]
FIELD_LABELS = {"customer_code": "کد مشتری (نرم‌افزار قبلی)", "receipt_no": "شماره فیش/پذیرش", "name": "نام مشتری", "first_name": "نام",
                "last_name": "نام خانوادگی", "mobile": "موبایل", "date": "تاریخ", "time": "ساعت", "appt_date": "تاریخ نوبت / مراجعه",
                "service": "خدمت", "line": "لاین", "staff": "پرسنل", "amount": "مبلغ", "status": "وضعیت", "settled_date": "تاریخ تسویه",
                "refund_date": "تاریخ استرداد", "notes": "توضیحات", "birth_date": "تاریخ تولد"}
# standard layouts: every file starts with the customer code, name and mobile, so all of them match up
TEMPLATES = {
    "customers": ["کد مشتری", "نام مشتری", "موبایل", "تاریخ تولد", "توضیحات"],
    "deposits": ["کد مشتری", "نام مشتری", "موبایل", "تاریخ پرداخت", "مبلغ بیعانه", "تاریخ مراجعه", "ساعت نوبت", "لاین", "خدمت",
                 "پرسنل", "وضعیت", "توضیحات"],
    "history": ["کد مشتری", "نام مشتری", "موبایل", "شماره فیش", "تاریخ", "ساعت", "خدمت", "لاین", "پرسنل", "مبلغ", "توضیحات"],
}
TEMPLATE_SAMPLES = {
    "customers": ["1001", "مریم احمدی", "09121234567", "1370/05/20", ""],
    "deposits": ["1001", "مریم احمدی", "09121234567", "1405/07/10", 500000, "1405/07/25", "11:30", "آرایش دائم", "فیبروز ابرو",
                 "لیلا تاجیک", "صندوق ودیعه", ""],
    "history": ["1001", "مریم احمدی", "09121234567", "1403060034", "1403/06/01", "19:15", "فیبروز ابرو", "آرایش دائم", "لیلا تاجیک",
                4500000, ""],
}


def template_xlsx(kind: str) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = KINDS[kind][:30]
    ws.sheet_view.rightToLeft = True
    ws.append(TEMPLATES[kind])
    ws.append(TEMPLATE_SAMPLES[kind])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="7C3AED")
        cell.alignment = Alignment(horizontal="center")
    for i, title in enumerate(TEMPLATES[kind], start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = max(12, len(title) + 6)
        for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
            row[0].number_format = "@" if title in ("کد مشتری", "موبایل", "شماره فیش") else row[0].number_format
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ImportProblem(ValueError):
    pass


def _key(text: object) -> str:
    """Comparable form of a header / name: Persian letters unified, no spaces or half-spaces."""
    return re.sub(r"[\s\-_‌.:]+", "", normalize_text(str(text or "")))


def _syn(kind: str | None) -> dict[str, list[str]]:
    merged = {**SYNONYMS, **KIND_SYNONYMS.get(kind or "", {})}
    if kind == "deposits":  # here «تاریخ مراجعه» is the appointment, never the payment date
        merged["date"] = [w for w in merged["date"] if w != "تاریخ مراجعه"]
    return {f: [_key(w) for w in words] for f, words in merged.items()}


_IGNORED = {_key(w) for w in IGNORED}


# ------------------------------------------------------------------ reading files
def read_table(filename: str, data: bytes) -> list[list[object]]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover - installed by requirements.txt
            raise ImportProblem("کتابخانه خواندن اکسل نصب نیست؛ update.bat را اجرا کنید") from exc
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        best: list[list[object]] = []
        for ws in wb.worksheets:  # the sheet with the most filled rows
            rows = [list(r) for r in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
            if len(rows) > len(best):
                best = rows
        wb.close()
        return best[:MAX_ROWS + 20]
    head = data[:2048].lstrip().lower()
    if head.startswith((b"<", b"\xef\xbb\xbf<")) or b"<table" in head or b"<html" in head:
        return _html_table(data)  # some programs save an HTML table with an .xls name
    if name.endswith(".xls") or data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        try:
            import xlrd
        except ImportError as exc:  # pragma: no cover - installed by requirements.txt
            raise ImportProblem("برای فایل xls، update.bat را اجرا کنید یا فایل را در Excel با «Save As» به .xlsx تبدیل کنید") from exc
        try:
            book = xlrd.open_workbook(file_contents=data)
        except Exception as exc:
            raise ImportProblem("این فایل xls خوانده نشد. آن را در Excel باز کنید و با «Save As» به صورت "
                                "«Excel Workbook (.xlsx)» ذخیره کنید، بعد همان را بارگذاری کنید.") from exc
        best: list[list[object]] = []
        for sh in book.sheets():
            rows = []
            for i in range(sh.nrows):
                vals = []
                for c in sh.row(i):
                    v = c.value
                    if c.ctype == xlrd.XL_CELL_DATE:
                        try:
                            v = xlrd.xldate.xldate_as_datetime(v, book.datemode)
                        except Exception:  # noqa: BLE001
                            pass
                    vals.append(v)
                if any(v not in (None, "") for v in vals):
                    rows.append(vals)
            if len(rows) > len(best):
                best = rows
        return best[:MAX_ROWS + 20]
    for enc in ("utf-8-sig", "cp1256", "utf-16"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ImportProblem("فایل قابل خواندن نیست؛ فایل Excel (.xlsx) یا CSV بفرستید")
    sample = text[:4000]
    delim = max([",", ";", "\t", "|"], key=sample.count)
    return [r for r in csv.reader(io.StringIO(text), delimiter=delim) if any(c.strip() for c in r)][:MAX_ROWS + 20]


def detect_mapping(header: list[object], kind: str | None = None) -> dict[str, int]:
    """field -> column index, from the column titles."""
    keys = [_key(h) for h in header]
    taken = {i for i, k in enumerate(keys) if k in _IGNORED}
    mapping: dict[str, int] = {}
    syn = _syn(kind)
    # exact titles first (in each field's order of preference), then titles that contain a known word
    for exact in (True, False):
        for field, words in syn.items():
            if field in mapping:
                continue
            for w in (words if exact else sorted(words, key=len, reverse=True)):
                hit = next((i for i, k in enumerate(keys) if k and i not in taken and i not in mapping.values()
                            and (k == w if exact else (len(w) >= 3 and w in k))), None)
                if hit is not None:
                    mapping[field] = hit
                    break
    if "first_name" in mapping and "last_name" not in mapping and "name" not in mapping:
        mapping["name"] = mapping.pop("first_name")
    return mapping


def _detect_line_column(rows: list[list[object]], mapping: dict[str, int]) -> int | None:
    """A column whose values look like «لاین ...» is the service line even if its title is something else."""
    width = max((len(r) for r in rows[:50]), default=0)
    for i in range(width):
        if i in mapping.values():
            continue
        vals = [_text(r[i]) for r in rows[:200] if i < len(r) and _text(r[i])]
        if len(vals) >= 3 and sum(v.startswith("لاین") for v in vals) / len(vals) > 0.6:
            return i
    return None


def _header_row(rows: list[list[object]], kind: str | None = None) -> int:
    scores = [len(detect_mapping(r, kind)) for r in rows[:15]]
    return max(range(len(scores)), key=lambda i: scores[i]) if scores else 0


def _html_table(data: bytes) -> list[list[object]]:
    from html.parser import HTMLParser

    text = data.decode("utf-8", errors="ignore") if b"\x00" not in data[:200] else data.decode("utf-16", errors="ignore")

    class P(HTMLParser):
        def __init__(self) -> None:
            super().__init__()
            self.rows: list[list[str]] = []
            self.row: list[str] | None = None
            self.cell: list[str] | None = None

        def handle_starttag(self, tag, attrs):  # noqa: ANN001
            if tag == "tr":
                self.row = []
            elif tag in ("td", "th") and self.row is not None:
                self.cell = []

        def handle_endtag(self, tag):  # noqa: ANN001
            if tag in ("td", "th") and self.row is not None and self.cell is not None:
                self.row.append(" ".join("".join(self.cell).split()))
                self.cell = None
            elif tag == "tr" and self.row is not None:
                if any(self.row):
                    self.rows.append(self.row)
                self.row = None

        def handle_data(self, d):  # noqa: ANN001
            if self.cell is not None:
                self.cell.append(d)

    p = P()
    p.feed(text)
    if not p.rows:
        raise ImportProblem("جدولی در فایل پیدا نشد")
    return p.rows[:MAX_ROWS + 20]


# ------------------------------------------------------------------ values
def parse_date(value: object) -> tuple[date | None, time | None]:
    if value in (None, ""):
        return None, None
    if isinstance(value, datetime):
        return value.date(), (value.time() if (value.hour or value.minute) else None)
    if isinstance(value, date):
        return value, None
    s = to_en_digits(str(value)).strip()
    t = None
    tm = re.search(r"(\d{1,2}):(\d{2})", s)
    if tm and 0 <= int(tm.group(1)) < 24 and 0 <= int(tm.group(2)) < 60:
        t = time(int(tm.group(1)), int(tm.group(2)))
        s = s.replace(tm.group(0), " ")
    m = re.search(r"(\d{1,4})\s*[/\-.]\s*(\d{1,2})\s*[/\-.]\s*(\d{1,4})", s)
    if m:
        a, b, c = (int(x) for x in m.groups())
        y, mo, d = (c, b, a) if c > 31 else (a, b, c)  # dd/mm/yyyy or yyyy/mm/dd
    else:
        m = re.search(r"\b(1[34]\d{2}|20\d{2})(\d{2})(\d{2})\b", s)  # 14030512
        if not m:
            return None, t
        y, mo, d = (int(x) for x in m.groups())
    if y < 100:
        y += 1400 if y < 50 else 1300
    try:
        if y < 1700:
            if not (1 <= mo <= 12 and 1 <= d <= 31):
                return None, t
            return date(*jalali_to_gregorian(y, mo, d)), t
        return date(y, mo, d), t
    except ValueError:
        return None, t


def parse_time(value: object) -> time | None:
    if isinstance(value, time):
        return value
    if isinstance(value, datetime):
        return value.time()
    _, t = parse_date(f"0 {value}") if value not in (None, "") else (None, None)
    if t is None and value not in (None, ""):
        m = re.fullmatch(r"\s*(\d{1,2})\s*", to_en_digits(str(value)))
        if m and int(m.group(1)) < 24:
            t = time(int(m.group(1)), 0)
    return t


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def _amount(value: object, unit: str) -> int | None:
    if value in (None, ""):
        return None
    n = int(value) if isinstance(value, (int, float)) else parse_amount(_text(value))
    if n is None:
        return None
    return int(n) * (10 if unit == "toman" else 1)


def _mobile(value: object) -> str | None:
    s = _text(value)
    if re.fullmatch(r"9\d{9}", s):  # Excel dropped the leading zero
        s = "0" + s
    return normalize_mobile(s)


# ------------------------------------------------------------------ lookups
class _Catalog:
    def __init__(self, db: Session):
        self.services: dict[str, int] = {}
        # active services first: when an archived and an active service share a name, the active one is used
        for s in db.scalars(select(Service).order_by(Service.is_active.desc(), Service.id)):
            for n in [s.name, *(s.aliases or [])]:
                self.services.setdefault(_key(n), s.id)
            if s.line:  # "line / service" or "line - service" as written by some exports
                self.services.setdefault(_key(f"{s.line.name}{s.name}"), s.id)
        self.lines: dict[str, int] = {}
        self.line_services: dict[int, list[int]] = {}
        for ln in db.scalars(select(ServiceLine)):
            self.lines.setdefault(_key(ln.name), ln.id)
            self.lines.setdefault(_key(re.sub(r"^\s*لاین\s*", "", ln.name)), ln.id)
        for sv in db.scalars(select(Service).where(Service.is_active.is_(True))):
            self.line_services.setdefault(sv.line_id, []).append(sv.id)
        self.staff_line: dict[int, int | None] = {}
        self.staff: dict[str, int] = {}
        for p in db.scalars(select(Staff)):
            self.staff_line[p.id] = p.line_id
            self.staff.setdefault(_key(p.full_name), p.id)
            first = p.full_name.split()[0] if p.full_name.split() else ""
            if first:
                self.staff.setdefault(_key(first), p.id)

    def line(self, name: str) -> int | None:
        if not name:
            return None
        return self.lines.get(_key(name)) or self.lines.get(_key(re.sub(r"^\s*لاین\s*", "", name)))

    def service(self, name: str) -> int | None:
        return self.services.get(_key(name)) if name else None

    def person(self, name: str) -> int | None:
        if not name:
            return None
        k = _key(name)
        return self.staff.get(k) or self.staff.get(_key(name.split()[0]))


# ------------------------------------------------------------------ preview
def parse_rows(db: Session, rows: list[list[object]], kind: str, unit: str, mapping: dict[str, int] | None = None) -> dict:
    if kind not in KINDS:
        raise ImportProblem("نوع اطلاعات نامعتبر است")
    if not rows:
        raise ImportProblem("فایل خالی است")
    hi = _header_row(rows, kind)
    header = [_text(h) for h in rows[hi]]
    mapping = {k: int(v) for k, v in (mapping or detect_mapping(rows[hi], kind)).items() if v is not None and int(v) >= 0}
    if "line" not in mapping:
        col = _detect_line_column(rows[hi + 1:], mapping)
        if col is not None:
            mapping["line"] = col
    if kind == "history" and "date" not in mapping and "appt_date" in mapping:
        mapping["date"] = mapping.pop("appt_date")
    need = {"customers": [], "history": ["date"], "deposits": ["amount"]}[kind]
    if not ({"name", "mobile", "first_name", "last_name", "customer_code"} & mapping.keys()):
        raise ImportProblem("ستون نام یا موبایل مشتری پیدا نشد؛ ستون‌ها را دستی مشخص کنید")
    missing = [FIELD_LABELS[f] for f in need if f not in mapping]
    if missing:
        raise ImportProblem(f"ستون «{'، '.join(missing)}» پیدا نشد؛ ستون‌ها را دستی مشخص کنید")

    cat = _Catalog(db)
    now = local_now()
    get = lambda r, f: r[mapping[f]] if f in mapping and mapping[f] < len(r) else None  # noqa: E731
    parsed: list[dict] = []
    skipped = {"settled": 0, "refunded": 0}
    receipt = ""  # grouped reports: a «شماره پذیرش: ...» line above the services of each receipt
    seen: dict[str, int] = {}
    for n, r in enumerate(rows[hi + 1: hi + 1 + MAX_ROWS], start=hi + 2):
        line_text = to_en_digits(" ".join(_text(c) for c in r if c not in (None, "")))
        group = re.search(r"(?:شماره\s*پذیرش|شماره\s*فیش|شماره\s*فاکتور)\s*[:：]?\s*(\d{3,})", line_text)
        if group and not _text(get(r, "service")) and not _text(get(r, "name")):
            receipt = group.group(1)
            continue
        name = _text(get(r, "name")) or " ".join(x for x in (_text(get(r, "first_name")), _text(get(r, "last_name"))) if x)
        if re.fullmatch(r"(جمع|جمع کل|total|مجموع).*", name or "", flags=re.I):
            continue
        raw_mobile = _text(get(r, "mobile"))
        code = _text(get(r, "customer_code"))
        row = {"row": n, "name": name, "mobile": _mobile(get(r, "mobile")), "mobile_raw": raw_mobile,
               "code": to_en_digits(code) if code not in ("0", "") else "",
               "errors": [], "warnings": []}
        if raw_mobile and not row["mobile"]:
            row["warnings"].append(f"موبایل «{raw_mobile}» ناقص یا نامعتبر است" + (" (مشتری با کد شناخته می‌شود)" if row["code"] else ""))
        if not row["name"] and not row["mobile"] and not row["code"]:
            # blank lines, row numbers only, or a totals line ("جمع کل") - not a record
            if any(re.match(r"\s*(جمع|مجموع|total)", _text(c), flags=re.I) for c in r) or \
                    not (_text(get(r, "date")) or _text(get(r, "service")) or _text(get(r, "appt_date"))):
                continue
            row["errors"].append("نام و موبایل مشتری خالی است")
        # deposits already used (settled) and refunded receipts are not brought over - but their customer is
        # (code, name, mobile), otherwise customers who only had settled deposits would lose their mobile
        status = _text(get(r, "status"))
        if kind == "deposits" and (re.search(r"تسویه|استفاده|مسترد|برگشت|باطل", status) or _text(get(r, "settled_date"))):
            skipped["settled"] += 1
            row["customer_only"] = True
        elif _text(get(r, "refund_date")) or re.search(r"مسترد|برگشت|باطل", status):
            skipped["refunded"] += 1
            row["customer_only"] = True
        d, t = parse_date(get(r, "date"))
        t = parse_time(get(r, "time")) or t
        if kind != "customers" and get(r, "date") not in (None, "") and d is None:
            row["errors"].append(f"تاریخ «{_text(get(r, 'date'))}» قابل خواندن نیست")
        row["date"] = datetime.combine(d, t or time(12, 0)).isoformat(timespec="minutes") if d else None
        ad, at_ = parse_date(get(r, "appt_date"))
        if kind == "deposits" and "time" in mapping and "appt_date" in mapping:
            at_ = parse_time(get(r, "time")) or at_
        row["appt_date"] = datetime.combine(ad, at_ or time(12, 0)).isoformat(timespec="minutes") if ad else None
        row["appt_time_known"] = bool(at_)
        row["service_name"] = _text(get(r, "service"))
        row["service_id"] = cat.service(row["service_name"])
        row["staff_name"] = _text(get(r, "staff"))
        row["staff_id"] = cat.person(row["staff_name"])
        row["line_name"] = _text(get(r, "line"))
        row["line_id"] = cat.line(row["line_name"]) or (cat.staff_line.get(row["staff_id"]) if row["staff_id"] else None)
        if row["line_name"] and not cat.line(row["line_name"]):
            row["warnings"].append(f"لاین «{row['line_name']}» در سیستم نیست")
        if not row["service_id"] and not row["service_name"] and row["line_id"] and len(cat.line_services.get(row["line_id"], [])) == 1:
            row["service_id"] = cat.line_services[row["line_id"]][0]  # the line has a single service
        if row["staff_name"] and not row["staff_id"]:
            row["warnings"].append(f"پرسنل «{row['staff_name']}» در سیستم نیست (در توضیحات ثبت می‌شود)")
        row["amount"] = _amount(get(r, "amount"), unit)
        row["notes"] = _text(get(r, "notes"))
        row["receipt"] = receipt or _text(get(r, "receipt_no"))
        bd, _ = parse_date(get(r, "birth_date"))
        row["birth_date"] = bd.isoformat() if bd else None
        if kind == "history":
            if not d:
                row["errors"].append("تاریخ خدمت خالی است")
            if not row["service_name"]:
                row["warnings"].append("نام خدمت خالی است")
            if d and d > now.date():
                row["warnings"].append("تاریخ آینده است؛ به‌صورت نوبت رزرو ثبت می‌شود")
        if kind == "deposits":
            if not row["amount"] or row["amount"] <= 0:
                row["errors"].append("مبلغ بیعانه خالی یا نامعتبر است")
            if d and d > now.date() and not ad:  # a single future date is the appointment, not the payment
                row["appt_date"], row["date"] = row["date"], None
                row["appt_time_known"] = bool(t)
            if row["appt_date"] and row["appt_date"][:10] < now.date().isoformat():
                row["warnings"].append("تاریخ مراجعه گذشته ولی بیعانه هنوز باز است؛ فقط بیعانه ثبت می‌شود")
            elif row["appt_date"] and not row["appt_time_known"]:
                row["warnings"].append("ساعت نوبت مشخص نیست؛ نوبت با «ساعت نامشخص» ثبت می‌شود")
        if row.get("customer_only"):  # only the customer's details are used from this line
            row["errors"] = [e for e in row["errors"] if "مشتری" in e]
            row["warnings"] = [w for w in row["warnings"] if "موبایل" in w]
            parsed.append(row)
            continue
        # fingerprint: the same record imported again (even from another file) is recognised and skipped
        who = row["code"] or row["mobile"] or _key(row["name"])
        base = "|".join(str(x) for x in (kind, who, row["receipt"] if kind == "history" else "", row["date"], row["service_name"],
                                         row["amount"], row["appt_date"] if kind == "deposits" else "", row["staff_name"]))
        seen[base] = seen.get(base, 0) + 1
        row["fp"] = f"{base}|{seen[base]}"
        parsed.append(row)

    unknown: dict[str, int] = {}
    unknown_line: dict[str, dict[int, int]] = {}
    unknown_staff: dict[str, int] = {}
    for p in parsed:
        if p["service_name"] and not p["service_id"] and not p["errors"] and not p.get("customer_only"):
            unknown[p["service_name"]] = unknown.get(p["service_name"], 0) + 1
            if p.get("line_id"):
                per = unknown_line.setdefault(p["service_name"], {})
                per[p["line_id"]] = per.get(p["line_id"], 0) + 1
        if p["staff_name"] and not p["staff_id"]:
            unknown_staff[p["staff_name"]] = unknown_staff.get(p["staff_name"], 0) + 1
    valid = [p for p in parsed if not p["errors"]]  # incl. lines used only for the customer's details
    ok = [p for p in valid if not p.get("customer_only")]
    # one mobile for several customers (in the file or already in the system): only a warning, the code decides
    by_mobile: dict[str, set[str]] = {}
    for p in valid:
        if p["mobile"]:
            by_mobile.setdefault(p["mobile"], set()).add(p["code"] or _key(p["name"]))
    owners = {m: (code, name) for m, code, name in db.execute(
        select(Customer.mobile, Customer.legacy_code, Customer.full_name).where(Customer.mobile.in_(list(by_mobile))))} if by_mobile else {}
    for p in valid:
        m = p["mobile"]
        if m and len(by_mobile[m]) > 1:
            p["warnings"].append(f"موبایل {m} در فایل برای چند مشتری آمده (تکراری)")
        elif m and m in owners and p["code"] and owners[m][0] != p["code"] and _key(owners[m][1]) != _key(p["name"]):
            p["warnings"].append(f"موبایل {m} در سیستم متعلق به «{owners[m][1]}» (کد {owners[m][0]}) است (تکراری)")
    people = {p["code"] or p["mobile"] or _key(p["name"]) for p in valid}
    codes = {p["code"] for p in valid if p["code"]}
    known = set(db.scalars(select(Customer.legacy_code).where(Customer.legacy_code.in_(codes)))) if codes else set()
    loose = {p["mobile"] for p in valid if p["mobile"] and not p["code"]}
    known_m = set(db.scalars(select(Customer.mobile).where(Customer.mobile.in_(loose)))) if loose else set()
    existing = {f"c:{c}" for c in known} | {f"m:{m}" for m in known_m}
    mobile_issues = sum(1 for p in valid if (p.get("mobile_raw") and not p["mobile"]) or any("تکراری" in w for w in p["warnings"]))
    dates = sorted(p["date"] for p in ok if p["date"])
    return {
        "kind": kind, "kind_label": KINDS[kind], "unit": unit, "header_row": hi + 1, "headers": header, "mapping": mapping,
        "fields": FIELD_LABELS, "rows": parsed,
        "summary": {"total": len(parsed), "ok": len(ok), "errors": len(parsed) - len(valid),
                    "warnings": sum(1 for p in ok if p["warnings"]), "customers": len(people),
                    "existing_customers": len(existing), "new_customers": max(0, len(people) - len(existing)),
                    "mobile_issues": mobile_issues,
                    "skipped_settled": skipped["settled"], "skipped_refunded": skipped["refunded"],
                    "amount": sum(p["amount"] or 0 for p in ok), "first_date": dates[0] if dates else None,
                    "last_date": dates[-1] if dates else None,
                    "future_appointments": sum(1 for p in ok if p.get("appt_date") and p["appt_date"] >= now.isoformat()[:16])},
        "unknown_services": sorted((_unknown_service(db, k, v, unknown_line.get(k)) for k, v in unknown.items()),
                                   key=lambda x: -x["count"]),
        "unknown_staff": sorted(({"name": k, "count": v} for k, v in unknown_staff.items()), key=lambda x: -x["count"]),
    }


def _unknown_service(db: Session, name: str, count: int, lines: dict[int, int] | None) -> dict:
    """A service name of the file that isn't in the system: the line the file gives it and look-alike services
    (a typo in the old software - «کروبکسی» for «کربوکسی» - should be mapped, not created twice)."""
    line_id = max(lines, key=lines.get) if lines else None
    line = db.get(ServiceLine, line_id) if line_id else None
    similar = [{"id": s.id, "name": s.name, "code": s.code, "line": s.line.name if s.line else None, "is_active": s.is_active}
               for s in service_catalog.similar(db, name)]
    return {"name": name, "count": count, "line_id": line_id, "line": line.name if line else None, "similar": similar}


def preview(db: Session, filename: str, data: bytes, kind: str, unit: str = "rial", mapping: dict | None = None) -> ImportBatch:
    result = parse_rows(db, read_table(filename, data), kind, unit, mapping)
    batch = ImportBatch(kind=f"legacy_{kind}", file_name=filename[:250], status="review", provider="legacy",
                        rows=result.pop("rows"), summary=result)
    db.add(batch)
    db.flush()
    return batch


# ------------------------------------------------------------------ commit
def commit(db: Session, batch: ImportBatch, *, service_map: dict[str, int | str] | None = None, new_line_id: int | None = None,
           payment_account_id: int | None = None, user=None) -> dict:  # noqa: ANN001
    if batch.status != "review":
        raise ImportProblem("این فایل قبلاً ثبت یا لغو شده است")
    kind = batch.kind.removeprefix("legacy_")
    now = local_now()
    service_map = service_map or {}
    created: dict[str, list[int]] = {"customers": [], "appointments": [], "deposits": [], "services": []}
    counts = {"customers_new": 0, "customers_updated": 0, "appointments": 0, "deposits": 0, "skipped_duplicates": 0, "skipped_errors": 0}

    # services that were not found: map to an existing one, create, or leave the name in the notes
    new_services: dict[str, int] = {}
    # a new service goes to the line the file gives it (most rows decide); without one, to the chosen line
    name_lines: dict[str, dict[int, int]] = {}
    for r in batch.rows:
        if r.get("service_name") and not r.get("service_id") and r.get("line_id"):
            per = name_lines.setdefault(service_catalog.name_key(r["service_name"]), {})
            per[r["line_id"]] = per.get(r["line_id"], 0) + 1

    def service_for(row: dict) -> int | None:
        if row.get("service_id"):
            return row["service_id"]
        name = row.get("service_name")
        if not name:
            return None
        choice = service_map.get(name, "new")
        if isinstance(choice, int) or (isinstance(choice, str) and choice.isdigit()):
            return int(choice)
        if choice != "new":
            return None
        key = service_catalog.name_key(name)
        if key not in new_services:
            per = name_lines.get(key)
            line_id = (max(per, key=per.get) if per else None) or new_line_id or _import_line(db).id
            svc = service_catalog.find_by_name(db, line_id, name)  # e.g. created a moment ago by another spelling
            line = db.get(ServiceLine, line_id)
            if line is not None and not line.is_active:
                line.is_active = True  # a service is about to be added to it: the line must show in the settings
            if svc is None:
                svc = Service(line_id=line_id, name=name.strip()[:120], base_price=0, duration_minutes=60, is_active=True)
                db.add(svc)
                db.flush()
                created["services"].append(svc.id)
            new_services[key] = svc.id
        return new_services[key]

    customers: dict[str, Customer] = {}

    # customer codes: the old software's code is kept; customers without one get the next free number after all
    # codes (in the database and in this file), so a code given here never collides with one coming later in the file
    file_codes = [int(r["code"]) for r in batch.rows if str(r.get("code") or "").isdigit()]
    auto_code = [max(int(accounting.next_customer_code(db)), max(file_codes, default=0) + 1)]
    moved: list[Customer] = []

    def new_code() -> str:
        auto_code[0] += 1
        return str(auto_code[0] - 1)

    def same_person(c: Customer, row: dict) -> bool:
        if row["mobile"] and c.mobile == row["mobile"]:
            return True
        return bool(row["name"]) and _key(c.full_name) == _key(row["name"])

    def mobile_state(row: dict, owner: Customer | None = None) -> tuple[str | None, str | None, str | None]:
        """(mobile to store, issue, raw number) - a wrong or already-used number is only flagged, never blocks."""
        m = row["mobile"]
        if m:
            other = db.scalar(select(Customer).where(Customer.mobile == m))
            if other is not None and other is not owner:
                return None, "duplicate", m
            return m, None, None
        if row.get("mobile_raw"):
            return None, "invalid", row["mobile_raw"]
        return None, "missing", None

    def customer_for(row: dict) -> Customer:
        code = row.get("code") or ""
        key = f"code:{code}" if code else (row["mobile"] or "name:" + _key(row["name"]))
        if key in customers:
            return customers[key]
        c = None
        if code:  # the customer code links the old software's reports (deposits, receipts, customers) together
            c = db.scalar(select(Customer).where(Customer.legacy_code == code))
            if c is None:  # the code of a duplicate record that was merged into another customer
                from .customer_merge import find_by_other_code
                c = find_by_other_code(db, code)
                merged_code = c is not None
            else:
                merged_code = False
            if c is not None and not merged_code and c.source != "import" and not same_person(c, row):
                # that code was given automatically in this system to someone else: the old code wins, theirs moves
                c.legacy_code = None
                moved.append(c)
                db.flush()
                counts["codes_moved"] = counts.get("codes_moved", 0) + 1
                c = None
            if c is None and row["mobile"]:  # same person already entered by hand in this system: give them the old code
                m = accounting.find_customer(db, row["mobile"])
                if m is not None and m.source != "import" and same_person(m, row):
                    c = m
                    if m.mobile == row["mobile"]:
                        c.legacy_code = code
                    else:  # found by the number of a merged duplicate: keep its own code, remember this one too
                        c.other_codes = list(dict.fromkeys([*(c.other_codes or []), code]))
        elif row["mobile"]:
            c = accounting.find_customer(db, row["mobile"])
        elif row["name"]:
            c = db.scalar(select(Customer).where(Customer.full_name == row["name"], Customer.mobile.is_(None)))
        if c is None:
            mobile, issue, raw = mobile_state(row)
            c = Customer(full_name=row["name"] or f"مشتری {code or row.get('mobile_raw') or ''}".strip(), mobile=mobile,
                         mobile_issue=issue, mobile_raw=raw, source="import", legacy_code=code or new_code())
            db.add(c)
            db.flush()
            created["customers"].append(c.id)
            counts["customers_new"] += 1
            if issue in ("invalid", "duplicate"):
                counts["mobile_issues"] = counts.get("mobile_issues", 0) + 1
        else:
            if row["name"] and (not c.full_name or c.full_name.startswith("مشتری ")):
                c.full_name = row["name"]
            if not c.mobile:
                mobile, issue, raw = mobile_state(row, c)
                if mobile:
                    c.mobile, c.mobile_issue, c.mobile_raw = mobile, None, None
                elif issue != "missing" and not c.mobile_issue:
                    c.mobile_issue, c.mobile_raw = issue, raw
            counts["customers_updated"] += 1
        if row.get("birth_date") and not c.birth_date:
            c.birth_date = row["birth_date"]
        if kind == "customers" and row.get("notes") and row["notes"] not in (c.notes or ""):
            c.notes = ((c.notes + "\n") if c.notes else "") + row["notes"]
        customers[key] = c
        return c

    def note(row: dict, svc_id: int | None) -> str:
        parts = [IMPORT_NOTE]
        if row.get("service_name") and not svc_id:
            parts.append(f"خدمت: {row['service_name']}")
        if row.get("staff_name") and not row.get("staff_id"):
            parts.append(f"پرسنل: {row['staff_name']}")
        if row.get("line_name") and not svc_id:
            parts.append(row["line_name"] if row["line_name"].startswith("لاین") else f"لاین: {row['line_name']}")
        if row.get("receipt"):
            parts.append(f"فیش {row['receipt']}")
        if row.get("notes"):
            parts.append(row["notes"])
        return " - ".join(parts)

    account = None
    if kind == "deposits":
        account = db.get(PaymentAccount, payment_account_id) if payment_account_id else db.scalar(
            select(PaymentAccount).where(PaymentAccount.is_active.is_(True)).order_by(PaymentAccount.id))
        if account is None:
            raise ImportProblem("ابتدا در تنظیمات یک حساب دریافت (کارتخوان/کارت/صندوق) تعریف کنید")

    # records already brought over by an earlier (not undone) import of the same kind
    done_fps: set[str] = set()
    for b in db.scalars(select(ImportBatch).where(ImportBatch.kind == batch.kind, ImportBatch.status == "committed")):
        done_fps.update((b.summary or {}).get("fps", []))
    fps: list[str] = []

    for row in batch.rows:
        if row.get("errors"):
            counts["skipped_errors"] += 1
            continue
        if row.get("customer_only"):  # settled deposit / refunded receipt: keep the customer's details only
            customer_for(row)
            continue
        if row.get("fp") in done_fps:
            counts["skipped_duplicates"] += 1
            continue
        if row.get("fp"):
            fps.append(row["fp"])
        c = customer_for(row)
        if kind == "customers":
            continue
        svc_id = service_for(row)
        if kind == "history":
            at = datetime.fromisoformat(row["date"])
            a = Appointment(customer_id=c.id, service_id=svc_id, staff_id=row.get("staff_id"), start_at=at,
                            status="done" if at <= now else "booked", quoted_price=row.get("amount") or 0, notes=note(row, svc_id))
            db.add(a)
            db.flush()
            created["appointments"].append(a.id)
            counts["appointments"] += 1
        elif kind == "deposits":
            received = datetime.fromisoformat(row["date"]) if row.get("date") else now
            received = min(received, now)
            appt_id = None
            unknown_time = not row.get("appt_time_known")
            if row.get("appt_date") and (row["appt_date"][:10] >= now.date().isoformat() if unknown_time
                                         else datetime.fromisoformat(row["appt_date"]) >= now):
                start = datetime.fromisoformat(row["appt_date"])
                if unknown_time:
                    start = datetime.combine(start.date(), scheduling._hours(db)[0])
                a = Appointment(customer_id=c.id, service_id=svc_id, staff_id=accounting.default_staff_id(db, svc_id, row.get("staff_id")),
                                start_at=start, status="booked", time_unknown=True if unknown_time else None,
                                quoted_price=(db.get(Service, svc_id).base_price if svc_id else 0), notes=note(row, svc_id))
                db.add(a)
                db.flush()
                appt_id = a.id
                created["appointments"].append(a.id)
                counts["appointments"] += 1
            svc = db.get(Service, svc_id) if svc_id else None
            dep = Deposit(customer_id=c.id, amount=row["amount"], payment_account_id=account.id, received_at=received,
                          service_id=svc_id, appointment_id=appt_id, source="import", notes=note(row, svc_id),
                          staff_id=accounting.default_staff_id(db, svc_id, row.get("staff_id")))
            db.add(dep)
            db.flush()
            # opening balance: the money was received in the old system, so cash is not counted again
            accounting.post(db, f"بیعانه انتقالی {c.full_name}", [
                accounting.Leg(accounting.account(db, accounting.OPENING), debit=dep.amount, customer_id=c.id),
                accounting.Leg(accounting.account(db, accounting.DEPOSITS), credit=dep.amount, customer_id=c.id,
                               line_id=svc.line_id if svc else row.get("line_id"), staff_id=dep.staff_id),
            ], "deposit", dep.id, at=received)
            created["deposits"].append(dep.id)
            counts["deposits"] += 1
    for c in moved:  # customers whose automatic code was taken by an imported one get a fresh code
        c.legacy_code = new_code()
    batch.status = "committed"
    # the parsed file is not needed any more (undo uses "created", re-import checks use "fps"): keep only the problem
    # lines for reference, so big imports don't bloat the database
    batch.rows = [r for r in batch.rows if r.get("errors") or r.get("warnings")][:1000]
    batch.summary = {**batch.summary, "created": created, "result": counts, "fps": fps, "committed_at": now.isoformat(timespec="minutes")}
    return counts


def _import_line(db: Session) -> ServiceLine:
    name = "خدمات انتقالی"
    line = db.scalar(select(ServiceLine).where(ServiceLine.name == name))
    if line is None:
        line = ServiceLine(name=name, color="#94a3b8")
        db.add(line)
        db.flush()
    return line


# ------------------------------------------------------------------ undo
def undo(db: Session, batch: ImportBatch) -> dict:
    if batch.status != "committed":
        raise ImportProblem("فقط فایل‌های ثبت‌شده قابل برگشت هستند")
    created = (batch.summary or {}).get("created", {})
    kept = {"deposits": 0, "appointments": 0, "customers": 0, "services": 0}
    removed = {"deposits": 0, "appointments": 0, "customers": 0, "services": 0}
    for did in created.get("deposits", []):
        d = db.get(Deposit, did)
        if d is None:
            continue
        if d.status != "held":  # already used on an invoice / refunded: leave it
            kept["deposits"] += 1
            continue
        for e in db.scalars(select(JournalEntry).where(JournalEntry.ref_type == "deposit", JournalEntry.ref_id == did)):
            db.delete(e)
        db.delete(d)
        removed["deposits"] += 1
    db.flush()
    for aid in created.get("appointments", []):
        a = db.get(Appointment, aid)
        if a is None:
            continue
        if a.invoice_id or db.scalar(select(func.count(Deposit.id)).where(Deposit.appointment_id == aid)):
            kept["appointments"] += 1
            continue
        db.delete(a)
        removed["appointments"] += 1
    db.flush()
    for sid in created.get("services", []):
        s = db.get(Service, sid)
        if s is None:
            continue
        used = any(db.scalar(select(func.count(m.id)).where(m.service_id == sid)) for m in (Appointment, Deposit, InvoiceItem))
        if used:
            kept["services"] += 1
            continue
        db.delete(s)
        removed["services"] += 1
    unused = []
    for cid in created.get("customers", []):
        c = db.get(Customer, cid)
        if c is None:
            continue
        used = any(db.scalar(select(func.count(m.id)).where(m.customer_id == cid)) for m in (Appointment, Deposit, Invoice, Payment, WaitlistEntry))
        if used:
            kept["customers"] += 1
            continue
        unused.append(c)
    removed["customers"] = accounting.delete_customers(db, unused)
    batch.status = "undone"
    batch.summary = {**batch.summary, "undo": {"removed": removed, "kept": kept}}
    return {"removed": removed, "kept": kept}
